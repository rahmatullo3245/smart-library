
from pathlib import Path

from docx import Document


def extract_docx_pages(file_path: str | Path) -> list[dict]:
    """
    DOCX matnini ajratadi.
    DOCX ichida PDF kabi ishonchli sahifa chegaralari bo'lmagani uchun
    hozircha barcha matn bitta mantiqiy sahifaga birlashtiriladi.
    """
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"DOCX topilmadi: {path}")

    document = Document(str(path))
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip().replace("\n", " ")
                for cell in row.cells
            ]
            if any(cells):
                parts.append(" | ".join(cells))

    full_text = "\n".join(parts).strip()

    if not full_text:
        return []

    return [{"page_number": 1, "text": full_text}]
